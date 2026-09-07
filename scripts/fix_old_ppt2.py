# -*- coding: utf-8 -*-
"""第二轮：概念/过程表述修正（对照论文与代码核验后重写），仍保留原风格。
直接修改《00_FAST-LIVO2 全景解析（原风格修订版）.pptx》。"""
from pptx import Presentation

F = "00_FAST-LIVO2 全景解析（原风格修订版）.pptx"
prs = Presentation(F)
S = prs.slides


def find(slide, sub):
    for sh in slide.shapes:
        if sh.has_text_frame and sub in sh.text_frame.text:
            return sh
    raise KeyError(sub)


def edit_lines(slide, anchor, mapping):
    """按行（段落）替换：段落文本包含 old_sub 则整行替换为 new_line，保留该段格式。"""
    sh = find(slide, anchor)
    hit = 0
    for para in sh.text_frame.paragraphs:
        full = "".join(r.text for r in para.runs)
        for old, new in mapping:
            if old in full:
                if para.runs:
                    para.runs[0].text = new
                    for r in para.runs[1:]:
                        r.text = ""
                else:
                    para.text = new
                hit += 1
                break
    if hit != len(mapping):
        print(f"WARN slide anchor={anchor[:20]} hit {hit}/{len(mapping)}")


def replace_block(slide, anchor, new_lines):
    """整块按行覆写（保留每段格式），行数对齐原段落数。"""
    sh = find(slide, anchor)
    paras = sh.text_frame.paragraphs
    for i, p in enumerate(paras):
        line = new_lines[i] if i < len(new_lines) else ""
        if p.runs:
            p.runs[0].text = line
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.text = line
    if len(new_lines) > len(paras):
        print(f"WARN overflow anchor={anchor[:20]}: {len(new_lines)}>{len(paras)}")


def sub_edit(slide, old, new):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if old in r.text:
                    r.text = r.text.replace(old, new)
                    return


# 6 页：法向细化表述（论文默认关闭，非“持续优化”）
sub_edit(S[5], "法向细化在独立线程持续优化。",
         "法向细化（可选模块，论文默认配置关闭）按式(14)用历史块光度误差精化平面法向。")

# 12 页：不是“找最近平面”，而是哈希定位所在体素（含邻域）
sub_edit(S[11], "每个原始点直接去找地图里最近的局部平面",
         "每个原始点经哈希表定位到所在体素（含邻域体素），用其中的局部平面")
sub_edit(S[11], "查最近平面算距离", "哈希查体素平面算距离")

# 16 页：R⁻¹ 公式补全点协方差项（对照 voxel_map.cpp:448）
edit_lines(S[15], "R^-1 = 1 / (0.001 +", [
    ("σ_l = J_nq", "σ_平面 = J_nq·Σ_{n,c}·J_nq^T（平面不确定度）"),
    ("R^-1 = 1 / (0.001 + σ_l)",
     "R^-1 = 1/(0.001 + σ_平面 + n^T·Σ_p·n)（voxel_map.cpp:448，含点噪声）"),
])

# 17 页：外点剔除机制按代码实际（3σ 马氏门限+按层下探），非“平面不参与匹配”
edit_lines(S[16], "用途2：外点剔除", [
    ("用途2", "用途2：外点剔除—— 逐点残差做 3σ 马氏门限检验，超限点下探八叉子层用更细平面重试，仍失败则丢弃（voxel_map.cpp:737）"),
])

# 21 页（idx20）：参考块更新的触发/评分机制按论文 V-D 式(12) 重写
replace_block(S[20], "会更新（论文Fig.S3）", [
    "会更新（论文 V-D，式12）",
    "评分 S = (1−ω₁)·平均NCC + ω₁·视角余弦",
    "· NCC 高：外观与多数历史块相似",
    "  （MVS 技巧，避开动态物体）",
    "· 视角余弦高：正对平面、纹理更清晰",
    "· 得分最高的块成为参考块",
])

# 23 页（idx22）：金字塔层数表述统一
sub_edit(S[22], "patch_pyrimid_level: 4（实际用3层金字塔）",
         "patch_pyrimid_level: 4（论文三层/代码四层）")

# 24 页（idx23）：仿射计算过程按代码实际重写（单应闭式+3点差分，非4角点最小二乘）
sub_edit(S[23], "Step 2 射线-平面求交", "Step 2 构造单应矩阵（式13）")
replace_block(S[23], "patch的4个角点像素(u,v)", [
    "取 patch 中心像素与 u、v 方向",
    "各偏移 4 像素的共 3 个像素，",
    "经 cam2world 反投影为参考帧",
    "的射线方向向量",
    "（注：不做 4 角点拟合）",
])
replace_block(S[23], "平面先验的核心！", [
    "平面先验的核心！",
    "H = R·(nᵀp·I − t·nᵀ)",
    "n = LiDAR 平面法向",
    "p = 视觉点（参考帧坐标）",
    "闭式计算，无需逐点求交深度",
])
sub_edit(S[23], "Step 3 位姿变换", "Step 3 投影到当前帧")
replace_block(S[23], "4个3D点从参考帧", [
    "3 个方向向量经 H 变换后，",
    "用 world2cam 投影到当前帧，",
    "得到 3 个新的像素位置；",
    "R,t 由 ESIKF 当前状态给出",
    "(参考帧→当前帧相对位姿)",
])
sub_edit(S[23], "Step 4 投影", "Step 4 差分得仿射")
replace_block(S[23], "3D点投影到当前帧图像：", [
    "A 列0 = (px_du − px) / 4",
    "A 列1 = (px_dv − px) / 4",
    "得到 2×2 仿射矩阵",
    "（不是 2×3，也不是最小二乘）",
    "透视畸变隐含在投影差异中",
])
sub_edit(S[23], "Step 5 拟合", "Step 5 扭曲采样")
replace_block(S[23], "4对对应点", [
    "warpAffine 用 A⁻¹ 在参考图上",
    "逐像素反算采样坐标，",
    "双线性插值取灰度，",
    "生成与当前帧视角对齐的 patch",
    "（vio.cpp:293）",
])
sub_edit(S[23], "拿到2×3仿射矩阵后", "拿到2×2仿射矩阵后")
sub_edit(S[23], "→ 每个角点做射线-平面求交，深度各不相同",
         "→ 单应矩阵含法向 n，透视形变被完整建模")

# 25 页（idx24）：页脚补金字塔层数说明
sub_edit(S[24], "逆组合公式见Section VII-B末尾",
         "逆组合公式见Section VII-B末尾; 金字塔论文三层/代码四层(avia.yaml)")

# 27 页（idx26）：状态维数与参考块评分表述
sub_edit(S[26], "解决：20维状态向量加1维τ（逆曝光时间）",
         "解决：19维状态向量本就含1维τ（逆曝光时间）")
sub_edit(S[26], "参考块更新：多观测时选平均光度误差最小的作参考；用NCC+视角余弦综合评分动态",
         "参考块更新：按式(12)评分（加权平均NCC+视角余弦，ω₁随法向不确定度自适应），得分最高的观测块作参考")

prs.save(F)
print("saved", F)
